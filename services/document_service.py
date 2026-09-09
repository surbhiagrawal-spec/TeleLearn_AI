import os
import sys
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config import Config
from database.db_init import get_db_connection
from services.rag_service import chunk_text, store_chunks

os.makedirs(Config.UPLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in Config.ALLOWED_EXTENSIONS

def secure_filename_custom(filename):
    """Simple filename sanitizer that keeps the extension."""
    name, ext = os.path.splitext(filename)
    safe_name = "".join(c for c in name if c.isalnum() or c in (' ', '-', '_')).rstrip()
    safe_name = safe_name.replace(' ', '_') or 'document'
    return f"{safe_name[:50]}{ext.lower()}"

def extract_text_from_pdf(filepath):
    try:
        import PyPDF2
        text_parts = []
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return '\n'.join(text_parts)
    except ImportError:
        return _fallback_pdf_read(filepath)
    except Exception as e:
        return f"Error reading PDF: {str(e)}"

def _fallback_pdf_read(filepath):
    """Very basic PDF text extraction without PyPDF2."""
    try:
        with open(filepath, 'rb') as f:
            content = f.read().decode('latin-1', errors='ignore')
        import re
        text = re.sub(r'[^\x20-\x7E\n\t]', ' ', content)
        text = re.sub(r'\s+', ' ', text)
        return text[:50000]
    except Exception as e:
        return f"Could not extract PDF text: {str(e)}"

def extract_text_from_docx(filepath):
    try:
        import docx
        doc = docx.Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)
        return '\n'.join(paragraphs)
    except ImportError:
        return "python-docx not installed. Run: pip install python-docx"
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"

def extract_text_from_txt(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    except Exception as e:
        return f"Error reading TXT: {str(e)}"

def extract_text(filepath, file_type):
    file_type = file_type.lower()
    if file_type == 'pdf':
        return extract_text_from_pdf(filepath)
    elif file_type == 'docx':
        return extract_text_from_docx(filepath)
    elif file_type == 'txt':
        return extract_text_from_txt(filepath)
    return "Unsupported file type."

def save_document(user_id, file_obj, original_name):
    """Save uploaded file, extract text, chunk it, store in DB. Returns (doc_id, error)."""
    if not allowed_file(original_name):
        return None, "File type not allowed. Use PDF, TXT, or DOCX."

    file_type = original_name.rsplit('.', 1)[1].lower()
    safe_name = secure_filename_custom(original_name)
    unique_name = f"{uuid.uuid4().hex[:8]}_{safe_name}"
    filepath = os.path.join(Config.UPLOAD_FOLDER, unique_name)

    try:
        file_obj.save(filepath)
    except Exception as e:
        return None, f"Failed to save file: {str(e)}"

    file_size = os.path.getsize(filepath)
    text = extract_text(filepath, file_type)

    if not text or len(text.strip()) < 50:
        os.remove(filepath)
        return None, "Could not extract meaningful text from the file. Please check the file and try again."

    chunks = chunk_text(text)
    if not chunks:
        os.remove(filepath)
        return None, "File has no usable text content."

    conn = get_db_connection()
    try:
        cursor = conn.execute(
            "INSERT INTO documents (user_id, filename, original_name, file_type, file_size) VALUES (?,?,?,?,?)",
            (user_id, unique_name, original_name, file_type, file_size)
        )
        doc_id = cursor.lastrowid
        conn.commit()
    except Exception as e:
        conn.close()
        os.remove(filepath)
        return None, f"Database error: {str(e)}"
    finally:
        conn.close()

    store_chunks(doc_id, chunks)
    return doc_id, None

def get_user_documents(user_id):
    conn = get_db_connection()
    try:
        docs = conn.execute(
            "SELECT * FROM documents WHERE user_id = ? ORDER BY uploaded_at DESC",
            (user_id,)
        ).fetchall()
        return [dict(d) for d in docs]
    finally:
        conn.close()

def delete_document(doc_id, user_id):
    conn = get_db_connection()
    try:
        doc = conn.execute(
            "SELECT * FROM documents WHERE id = ? AND user_id = ?", (doc_id, user_id)
        ).fetchone()
        if not doc:
            return False, "Document not found."

        filepath = os.path.join(Config.UPLOAD_FOLDER, doc['filename'])
        if os.path.exists(filepath):
            os.remove(filepath)

        conn.execute("DELETE FROM documents WHERE id = ?", (doc_id,))
        conn.commit()
        return True, None
    except Exception as e:
        return False, str(e)
    finally:
        conn.close()

def search_documents(user_id, query):
    conn = get_db_connection()
    try:
        docs = conn.execute(
            """SELECT d.*, COUNT(dc.id) as chunk_count
               FROM documents d
               LEFT JOIN document_chunks dc ON d.id = dc.document_id
               WHERE d.user_id = ? AND (d.original_name LIKE ? OR dc.content LIKE ?)
               GROUP BY d.id
               ORDER BY d.uploaded_at DESC""",
            (user_id, f'%{query}%', f'%{query}%')
        ).fetchall()
        return [dict(d) for d in docs]
    finally:
        conn.close()
